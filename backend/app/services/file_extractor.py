import io
from typing import Optional
from fastapi import UploadFile, HTTPException
from pypdf import PdfReader
from docx import Document

class FileExtractor:
    @staticmethod
    async def extract_text_from_upload(file: UploadFile) -> str:
        filename = (file.filename or "").lower()
        contents = await file.read()
        
        if not contents:
            raise HTTPException(status_code=400, detail="Uploaded file is empty")

        if filename.endswith(".pdf"):
            return FileExtractor.extract_from_pdf(contents)
        elif filename.endswith(".docx"):
            return FileExtractor.extract_from_docx(contents)
        elif filename.endswith(".txt"):
            return FileExtractor.extract_from_txt(contents)
        else:
            raise HTTPException(
                status_code=400,
                detail="Unsupported file format. Please upload PDF, DOCX, or TXT."
            )

    @staticmethod
    def extract_from_pdf(file_bytes: bytes) -> str:
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            text_parts = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text.strip())
            full_text = "\n\n".join(text_parts).strip()
            if not full_text:
                raise ValueError("No extractable text found in PDF.")
            return full_text
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to read PDF resume: {str(e)}")

    @staticmethod
    def extract_from_docx(file_bytes: bytes) -> str:
        try:
            doc = Document(io.BytesIO(file_bytes))
            text_parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    text_parts.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text_parts.append(" | ".join(row_text))
            full_text = "\n".join(text_parts).strip()
            if not full_text:
                raise ValueError("No extractable text found in DOCX.")
            return full_text
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Failed to read DOCX resume: {str(e)}")

    @staticmethod
    def extract_from_txt(file_bytes: bytes) -> str:
        try:
            return file_bytes.decode("utf-8").strip()
        except UnicodeDecodeError:
            try:
                return file_bytes.decode("latin-1").strip()
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Failed to decode text file: {str(e)}")
